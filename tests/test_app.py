"""端到端测试：真浏览器打开网页，连本地的假 GitHub（tests/fake_github.py），把主要功能走一遍。

    pip install playwright openpyxl && python -m playwright install chromium
    python tests/test_app.py            # 全部
    python tests/test_app.py 借出 出差    # 只跑名字里含这些字的步骤（前面的「连接」总会跑）

天气（Open-Meteo）和 DeepSeek 用假数据代替，不联网、不花钱。失败时截图在 tests/artifacts/。
"""

import functools
import json
import re
import os
import struct
import subprocess
import sys
import threading
import traceback
import zlib
from datetime import date, timedelta
from functools import partial
from urllib.parse import quote
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from playwright.sync_api import expect, sync_playwright

sys.path.insert(0, str(Path(__file__).parent))
from fake_github import FakeRepo, serve  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
ART = ROOT / "tests" / "artifacts"
APP_PORT, API_PORT = 8765, 8766
URL = f"http://127.0.0.1:{APP_PORT}/"
API = f"http://127.0.0.1:{API_PORT}"
REPO = "test/inventory-data"
D = lambda n: (date.today() + timedelta(days=n)).isoformat()  # noqa: E731


# ---------- 测试数据 ----------

def seed():
    L = lambda i, name, parent=None: {"id": i, "name": name, "parent": parent, "assetId": None}  # noqa: E731
    tags = ["电子产品", "衣服", "运动服", "鞋", "包", "洗漱护肤", "清洁用品", "床上用品", "文具", "书籍资料", "证件文件",
            "钥匙", "水杯餐具", "收藏纪念", "旅行用品", "搬家用品", "收纳容器", "日用杂物", "零食食品", "药品急救", "运动器材"]
    locations = [L("Lmain", "主屋 Main Room"), L("Lward", "当季衣柜 Current Season Wardrobe", "Lmain"),
                 L("Lsnack", "零食食品柜 Snacks & Food Storage", "Lmain"), L("Ldesk", "书桌 Desk", "Lmain"),
                 L("Ldrawer", "书桌抽屉 Desk Drawer", "Ldesk"), L("Lbulk", "囤货柜 Bulk Items Backstock", "Lmain"),
                 L("Lshoe", "鞋柜 Shoe Cabinet", "Lmain"), L("Lstore", "储物间 Storage Room")]

    def I(i, name, tag, loc, asset=None, **kw):  # noqa: E743
        return {"id": i, "name": name, "assetId": asset, "location": loc, "tags": [tag], "quantity": kw.pop("quantity", 1),
                "description": "", "fields": kw.pop("fields", {}), "photos": [], "receipts": [], "notes": "", "archived": False,
                "createdAt": "2026-10-01T00:00:00Z", "updatedAt": "2026-10-01T00:00:00Z", **kw}

    items = [
        I("iw1", "黑色羽绒服", "衣服", "Lward", "110-001", fields={"季节": "冬", "颜色": "黑色", "部位": "外套", "厚薄": "厚"}, label="none"),
        I("iw2", "白色T恤", "衣服", "Lward", "110-002", fields={"季节": "夏", "颜色": "白色", "部位": "上衣", "厚薄": "薄"}, label="none"),
        I("iw3", "灰色卫衣", "衣服", "Lward", "110-003", fields={"季节": "春秋", "颜色": "灰色", "部位": "上衣", "厚薄": "适中"}, label="none"),
        I("ish", "白色运动鞋", "鞋", "Lshoe", "130-001", label="none"),
        I("iid", "身份证", "证件文件", "Ldrawer", "200-001", labelPrinted=False),
        I("ich", "手机充电器", "电子产品", "Ldrawer", "100-001", labelPrinted=False),
        I("itb", "牙刷", "洗漱护肤", "Lstore", "150-001", labelPrinted=False),
        I("imed", "布洛芬片", "药品急救", "Lbulk", "290-001", quantity=2, fields={"保质期": "2027-11-20"}, labelPrinted=False),
        I("ibook", "线性代数（第六版）", "书籍资料", "Ldrawer", "190-001", labelPrinted=False),
        I("isock", "黑袜子", "衣服", "Lward", "110-005", quantity=6, fields={"部位": "袜子"}, label="none",
          laundry={"state": "washing", "since": D(-1), "autoReturn": D(0)}),
        I("isheet", "灰色床单", "床上用品", "Lstore", "170-001", label="none", lastWashed=D(-20)),
        I("icold1", "999感冒灵颗粒", "药品急救", "Lbulk", "290-002", label="none"),
        I("icold2", "感康（复方氨酚烷胺片）", "药品急救", "Lbulk", "290-003", label="none"),
    ]
    locations.append(L("Lsummer", "夏季衣物与运动服柜 Summer Clothes & Sportswear Storage", "Lmain"))
    data = {"version": 1, "tags": tags, "locations": locations, "items": items, "prefs": {"homeCity": "济南"},
            "fieldPresets": {"衣服": ["季节", "颜色", "尺码"], "零食食品": ["保质期"], "药品急救": ["保质期", "规格", "剩余"]}}
    return {"inventory.json": json.dumps(data, ensure_ascii=False).encode()}


# 账本（顺手记账写进去）：只放用得到的，账户名、数字都是编的
def ledger_seed():
    cat = lambda i, name, group: {"id": i, "name": name, "kind": "expense", "group": group}  # noqa: E731
    data = {"version": 1, "accounts": [{"id": "a-live", "name": "测试卡", "currency": "CNY", "opening": 0},
                                       {"id": "a-usd", "name": "美元账户", "currency": "USD", "opening": 0}],
            "categories": [cat("c-snack", "零食", "food"), cat("c-drink", "饮料奶茶", "food"), cat("c-tissue", "纸巾清洁", "daily"),
                           cat("c-medical", "买药", "daily"), cat("c-gadget", "电子耗材", "daily"), cat("c-dorm", "宿舍小物件", "daily"),
                           cat("c-wish", "心愿", "none"), {"id": "i-job", "name": "兼职", "kind": "income"}],
            "budget": {"food": 1000, "daily": 500}, "settings": {"periodStartDay": 1},
            "tx": []}
    return {"finance.json": json.dumps(data, ensure_ascii=False).encode()}


def png(path, rgb=(200, 80, 60)):
    w = h = 64
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))
    chunk = lambda t, d: struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d))  # noqa: E731
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))
    return path


# ---------- 运行框架 ----------

STEPS = []


def step(name):
    def wrap(fn):
        STEPS.append((name, fn))
        return fn
    return wrap


class Ctx:
    def __init__(self, page, repo, ledger=None):
        self.page, self.repo, self.ledger = page, repo, ledger
        self.prompt = ""

    def ledger_tx(self):
        return json.loads(self.ledger.read("finance.json"))["tx"]

    def data(self):
        # 修改是先存手机、后台上传的：等待上传队列清空再读仓库
        self.page.wait_for_function(
            "() => { try { const q = JSON.parse(localStorage.getItem('inventory-queue')); return !q || !q.items.length; } catch { return true; } }",
            timeout=15000)
        return json.loads(self.repo.read("inventory.json"))

    def item(self, name):
        return next(i for i in self.data()["items"] if i["name"] == name)

    def go(self, hash_):
        # 地址没变时浏览器不会重新渲染（比如连着两次进「新建」），这时强制刷新
        same = self.page.url == URL + hash_
        self.page.goto(URL + hash_)
        if same:
            self.page.reload()

    def wait_item(self, name):
        self.page.wait_for_function("location.hash.startsWith('#/item/') && !location.hash.endsWith('/edit')", timeout=20000)
        expect(self.page.locator(".item-title")).to_contain_text(name)

    def new_item(self, name, tag, loc_label, **opts):
        p = self.page
        self.go("#/new")
        p.get_by_label("名称").fill(name)
        p.get_by_label("位置").select_option(label=loc_label)
        p.get_by_role("button", name=tag, exact=True).click()
        if opts.get("photo"):
            p.locator("section").first.locator("input[type=file]").set_input_files(str(opts["photo"]))
        p.get_by_role("button", name="保存", exact=True).click()
        self.wait_item(name)
        return self.item(name)


# ---------- 测试步骤 ----------

@step("连接：设置令牌后进入首页，旧数据自动转换")
def _(c):
    p = c.page
    c.go("#/settings")
    p.evaluate(f"() => {{ localStorage.clear(); localStorage.setItem('inventory-api-base', '{API}'); localStorage.setItem('inventory-test-scan', '1'); }}")
    p.reload()
    p.get_by_label("数据仓库").fill(REPO)
    p.get_by_role("textbox", name="令牌", exact=True).fill("test-token")
    p.get_by_label("令牌到期日").fill(D(5))
    p.get_by_role("button", name="保存并连接").click()
    expect(p.locator(".today-head")).to_be_visible()
    # 和「生活」一样的外观：问候、节气小标签、角落一句话
    expect(p.locator(".today-head .greet")).to_have_text(re.compile("好|夜深"))
    expect(p.locator(".head-tags .tag").first).to_have_text(re.compile("时节|还有|今天"))
    expect(p.locator(".whisper")).to_have_count(1)
    expect(p.get_by_text("GitHub 令牌还有")).to_be_visible()
    c.go("#/items")
    expect(p.get_by_text("13 件")).to_be_visible()
    expect(p.locator(".whisper")).to_have_count(0)   # 角落那句话只在首页
    expect(p.locator(".tile", has_text="手机充电器").locator(".tile-ph.tone-accent")).to_be_visible()   # 没照片：按类别的小图标
    expect(p.locator(".tile", has_text="布洛芬片").locator(".asset")).to_have_text("290-0001")
    p.get_by_role("button", name="列表").click()
    expect(p.locator(".row", has_text="布洛芬片")).to_be_visible()
    p.get_by_role("button", name="按位置").click()
    expect(p.locator(".row.place", has_text="囤货柜 Bulk Items Backstock")).to_be_visible()
    p.get_by_role("button", name="按类别").click()


@step("新建：按类别自动编号（4 位）、贴标签开关、照片、表单不被后台刷新清空")
def _(c):
    p = c.page
    c.go("#/new")
    expect(p.get_by_text("选好标签后会按类别自动编号")).to_be_visible()
    p.get_by_role("button", name="电子产品", exact=True).click()
    expect(p.get_by_role("textbox", name="编号")).to_have_value("100-0002")
    expect(p.locator(".switch-row input").first).to_be_checked()
    p.get_by_label("名称").fill("降噪耳机")
    p.get_by_label("位置").select_option(label="　　书桌抽屉 Desk Drawer")
    p.locator("section").first.locator("input[type=file]").set_input_files(str(png(ART / "p.png")))
    d = c.data(); d["tags"].append("外部标签"); c.repo.external_write("inventory.json", json.dumps(d, ensure_ascii=False).encode())
    p.evaluate("() => document.dispatchEvent(new Event('visibilitychange'))")
    p.wait_for_timeout(1500)
    expect(p.get_by_label("名称")).to_have_value("降噪耳机")
    expect(p.locator(".photo-cell")).to_have_count(1)
    p.get_by_role("button", name="保存", exact=True).click()
    c.wait_item("降噪耳机")
    it = c.item("降噪耳机")
    assert it["assetId"] == "100-0002" and it["label"] == "pending" and len(it["photos"]) == 1, it
    assert c.repo.read(it["photos"][0]["thumb"]), "缩略图没上传"
    assert "外部标签" in c.data()["tags"], "覆盖了别处的修改"
    c.go("#/new")
    p.get_by_role("button", name="衣服", exact=True).click()
    expect(p.get_by_role("textbox", name="编号")).to_have_value("110-0006")
    expect(p.locator(".switch-row input").first).not_to_be_checked()
    # 只能选一个类别：再点别的就换成别的
    p.get_by_role("button", name="文具", exact=True).click()
    expect(p.locator(".chips .chip.on")).to_have_count(1)
    expect(p.locator(".chips .chip.on")).to_have_text("文具")
    expect(p.get_by_role("textbox", name="编号")).to_have_value("180-0001")


@step("别的设备抢先用了推荐的编号：保存时自动顺延")
def _(c):
    p = c.page
    c.go("#/new")
    p.get_by_label("名称").fill("充电宝")
    p.get_by_label("位置").select_option(label="　　书桌抽屉 Desk Drawer")
    p.get_by_role("button", name="电子产品", exact=True).click()
    expect(p.get_by_role("textbox", name="编号")).to_have_value("100-0003")
    d = c.data()
    d["items"].append({**c.item("降噪耳机"), "id": "iext", "name": "外部物品", "assetId": "100-0003", "photos": []})
    c.repo.external_write("inventory.json", json.dumps(d, ensure_ascii=False).encode())
    p.get_by_role("button", name="保存", exact=True).click()
    c.wait_item("充电宝")
    assert c.item("充电宝")["assetId"] == "100-0004"


@step("标签：导出 Excel、标记已打印、重新打印")
def _(c):
    import openpyxl
    p = c.page
    c.go("#/labels")
    d = c.data()
    n = sum(1 for x in d["items"] + d["locations"] if x.get("assetId") and x.get("label") == "pending" and not x.get("archived"))
    expect(p.get_by_role("button", name=f"待打印（{n}）")).to_be_visible()
    with p.expect_download() as dl:
        p.get_by_role("button", name="下载 Excel").click()
    path = ART / "labels.xlsx"
    dl.value.save_as(path)
    rows = [[x.value for x in r] for r in openpyxl.load_workbook(path).active.iter_rows()]
    assert rows[0] == ["编号", "名称", "二维码"] and len(rows) == n + 1, rows
    p.get_by_role("button", name="标记为已打印").click()
    expect(p.get_by_role("button", name=f"已打印（{n}）")).to_be_visible()
    c.go(f"#/item/{c.item('身份证')['id']}")
    p.get_by_role("button", name="重新打印").click()
    expect(p.locator(".row-meta .label-state").first).to_have_text("待打印")


@step("消耗品：用掉一个、用完了、补货（编号不变）")
def _(c):
    p = c.page
    med = c.item("布洛芬片")
    c.go(f"#/item/{med['id']}")
    p.get_by_role("button", name="用掉一个").click()
    expect(p.get_by_role("button", name="用掉一个")).to_have_count(0)
    p.get_by_role("button", name="用完了", exact=True).click()
    expect(p.get_by_text("已用完，在购物清单上")).to_be_visible()
    c.go("#/me")
    expect(p.locator(".cell", has_text="购物清单")).to_contain_text("1")
    c.go(f"#/item/{med['id']}")
    p.get_by_role("button", name="补货").click()
    p.locator(".sheet input[type=number]").fill("3")
    p.locator(".sheet input[placeholder^='例如']").fill("2028-06")
    p.locator(".sheet").get_by_role("button", name="补货").click()
    expect(p.get_by_text("已用完，在购物清单上")).to_have_count(0)
    m = c.item("布洛芬片")
    assert m["quantity"] == 3 and m["fields"]["保质期"] == "2028-06" and m["assetId"] == "290-0001", m


@step("归档点选原因；删除后编号不复用")
def _(c):
    p = c.page
    c.go(f"#/item/{c.item('白色T恤')['id']}")
    p.get_by_role("button", name="归档").click()
    p.locator(".sheet").get_by_role("button", name="送人").click()
    p.locator(".sheet").get_by_role("button", name="归档").click()
    expect(p.get_by_text("已归档：送人")).to_be_visible()
    it = c.new_item("测试剪刀", "日用杂物", "　书桌 Desk")
    p.get_by_role("button", name="删除").click()
    p.wait_for_function("location.hash === '#/'")
    it2 = c.new_item("测试剪刀2", "日用杂物", "　书桌 Desk")
    assert (it["assetId"], it2["assetId"]) == ("270-0001", "270-0002"), (it["assetId"], it2["assetId"])


@step("借阅：借一本书、快到期提醒、续借、归还后归档")
def _(c):
    p = c.page
    c.go("#/borrow")
    p.get_by_role("button", name="借一本书").click()
    sh = p.locator(".sheet")
    sh.get_by_label("书名").fill("深度学习")
    sh.get_by_label("从哪借的").fill("学校图书馆")
    sh.get_by_label("借阅日期").fill(D(-28))
    sh.get_by_label("应还日期").fill(D(2))
    sh.get_by_role("button", name="记下").click()
    expect(p.locator(".row", has_text="深度学习")).to_contain_text("剩 2 天")
    book = c.item("深度学习")
    assert book["borrow"]["from"] == "学校图书馆" and book["tags"] == ["书籍资料"] and book["label"] == "none", book
    c.go("#/")
    expect(p.locator(".cell", has_text="1 本书要还了")).to_be_visible()
    c.go(f"#/item/{book['id']}")
    p.get_by_role("button", name="续借").click()
    p.locator(".sheet").get_by_label("新的应还日期").fill(D(30))
    p.locator(".sheet").get_by_role("button", name="续借").click()
    expect(p.get_by_text("还剩 30 天")).to_be_visible()
    p.get_by_role("button", name="已归还").click()
    expect(p.get_by_text("已归档：已归还")).to_be_visible()
    b = c.item("深度学习")
    assert b["archived"] and b["borrow"]["renewals"] == 1, b
    c.go("#/borrow")
    expect(p.locator(".section-title", has_text="已归还")).to_be_visible()


@step("装箱：新建箱子、装进去、全部放回原处")
def _(c):
    p = c.page
    c.go("#/boxes")
    p.get_by_role("button", name="新建搬家箱子").click()
    p.locator(".sheet").get_by_role("button", name="新建").click()
    expect(p.get_by_role("heading", name="箱子 1")).to_be_visible()
    box = next(l for l in c.data()["locations"] if l.get("box") == "move")
    assert box["assetId"].startswith("010-") and box["label"] == "pending", box
    charger = c.item("手机充电器")
    c.go(f"#/item/{charger['id']}/edit")
    p.get_by_label("位置").select_option(label="箱子 1")
    p.get_by_role("button", name="保存", exact=True).click()
    c.wait_item("手机充电器")
    expect(p.get_by_text("原来在 主屋 Main Room / 书桌 Desk / 书桌抽屉 Desk Drawer")).to_be_visible()
    c.go(f"#/place/{box['id']}")
    p.get_by_role("button", name="全部放回原处（1）").click()
    expect(p.get_by_text("0 件")).to_be_visible()
    it = c.item("手机充电器")
    assert it["location"] == "Ldrawer" and "homeLocation" not in it, it


@step("出差：天气 + DeepSeek 推荐和穿搭、装进行李箱、回来放回原处")
def _(c):
    p = c.page
    # 旧版存在本机的密钥：打开网页时自动搬进数据仓库的 config/ai.json
    p.evaluate("() => localStorage.setItem('inventory-deepseek', JSON.stringify({ key: 'sk-test', model: 'deepseek-chat' }))")
    c.go("#/")
    p.reload()
    expect(p.locator(".toast", has_text="DeepSeek 密钥存到数据仓库")).to_be_visible()
    assert json.loads(c.repo.read("config/ai.json"))["deepseek"]["key"] == "sk-test"
    assert p.evaluate("() => localStorage.getItem('inventory-deepseek')") is None
    c.go("#/settings")
    expect(p.get_by_text("已连接 DeepSeek")).to_be_visible()
    expect(p.get_by_role("textbox", name="API 密钥")).to_be_hidden()   # 已连接时收起
    assert "sk-test" not in c.repo.read("inventory.json").decode()
    c.go("#/trips")
    p.get_by_role("button", name="新出行").click()
    p.get_by_label("目的地").fill("上海")
    p.get_by_label("出发").fill(D(2))
    p.get_by_label("返回").fill(D(4))
    p.get_by_role("button", name="见客户").click()
    p.get_by_role("button", name="生成推荐").click()
    expect(p.get_by_text("由 DeepSeek 推荐和搭配")).to_be_visible(timeout=20000)
    expect(p.get_by_text("每天穿搭")).to_be_visible()
    expect(p.get_by_text("转换插头")).to_be_visible()
    expect(p.locator(".check-row", has_text="不存在的东西")).to_have_count(0)   # AI 编的 id 被过滤
    p.screenshot(path=ART / "trip.png", full_page=True)
    p.get_by_role("button", name="不核对，直接出发").click()
    expect(p.get_by_role("button", name="不核对，直接放回原处")).to_be_visible()
    d = c.data()
    trip = d["trips"][0]
    assert trip["status"] == "packed" and sorted(trip["checked"]) == ["iid", "iw1", "iw3"], trip
    assert all(next(i for i in d["items"] if i["id"] == x)["location"] == trip["boxId"] for x in trip["checked"])
    c.go("#/")
    expect(p.locator(".cell", has_text="行李箱里还有东西")).to_contain_text("上海")
    c.go(f"#/trip/{trip['id']}")
    p.get_by_role("button", name="不核对，直接放回原处").click()
    expect(p.get_by_role("button", name="再来一次")).to_be_visible()
    d = c.data()
    assert next(i for i in d["items"] if i["id"] == "iw1")["location"] == "Lward"
    assert not any(l["id"] == trip["boxId"] for l in d["locations"]), "空行李箱没删掉"


@step("出差：没有 DeepSeek 时用规则推荐")
def _(c):
    p = c.page
    c.go("#/settings")
    p.get_by_text("更换或删除密钥").click()
    p.get_by_role("textbox", name="API 密钥").fill("")
    p.get_by_role("button", name="测试并保存到数据仓库").click()
    expect(p.get_by_text("还没有填")).to_be_visible()
    assert "deepseek" not in json.loads(c.repo.read("config/ai.json"))
    c.go("#/trips")
    p.get_by_role("button", name="新出行").click()
    p.get_by_label("目的地").fill("哈尔滨")
    p.get_by_label("出发").fill(D(1))
    p.get_by_label("返回").fill(D(3))
    p.get_by_role("button", name="生成推荐").click()
    expect(p.get_by_text("规则推荐")).to_be_visible(timeout=20000)
    for name in ["黑色羽绒服", "身份证", "手机充电器", "布洛芬片"]:
        expect(p.locator(".check-row", has_text=name)).to_be_visible()
    # 两种感冒药只带一种
    expect(p.locator(".check-row", has_text="感冒灵").or_(p.locator(".check-row", has_text="感康"))).to_have_count(1)
    expect(p.locator(".check-row", has_text="白色T恤")).to_have_count(0)   # 已归档，且是夏装


@step("今天穿什么：首页推荐、换一套、自己点选今天穿的、白天改一下")
def _(c):
    p = c.page
    p.evaluate("() => localStorage.setItem('inventory-deepseek', JSON.stringify({ key: 'sk-test' }))")
    c.go("#/")
    p.reload()
    card = p.locator(".outfit-card")
    # 首页第一次打开时（还没有密钥）已经按规则生成过今天的搭配；有了 DeepSeek 后重新推荐
    expect(card.locator("h2")).to_be_visible(timeout=20000)
    c.go("#/outfit")
    p.get_by_text("换个安排重新推荐").click()
    p.get_by_role("button", name="宅宿舍").click()
    p.get_by_role("button", name="见客户").click()
    p.get_by_role("button", name="重新推荐", exact=True).click()
    expect(p.get_by_text("由 DeepSeek 搭配")).to_be_visible(timeout=20000)
    assert "见客户" in c.data()["outfit"]["schedule"]
    c.go("#/")
    expect(card.locator("h2")).to_have_text("灰卫衣配黑羽绒", timeout=20000)
    p.get_by_role("button", name="换一套（1/2）").click()
    expect(card.locator("h2")).to_have_text("暖和的一套")
    # 推荐只打标记，不预选；自己一件件点
    card.get_by_role("link", name="选今天穿的").click()
    pick = p.locator("#pick")
    expect(pick.locator(".garment.selected")).to_have_count(0)
    expect(pick.locator(".garment", has_text="灰色卫衣").locator(".rec")).to_be_visible()
    for name in ["灰色卫衣", "黑色羽绒服", "白色运动鞋"]:
        pick.locator(".garment", has_text=name).click()
    p.get_by_role("button", name="就穿这些（3 件）").click()
    p.wait_for_function("!location.hash.startsWith('#/outfit')", timeout=20000)   # 保存完会离开挑选页
    worn = {i["id"]: i.get("wearsSinceWash") for i in c.data()["items"] if i.get("worn")}
    assert worn == {"iw3": 1, "iw1": 1, "ish": 1}, worn
    # 白天改一下：鞋换掉，次数跟着更正
    c.go("#/")
    expect(card.locator("h2")).to_have_text("今天穿的")
    card.get_by_role("link", name="改一下").click()
    expect(pick.locator(".garment.selected")).to_have_count(3)
    pick.locator(".garment", has_text="白色运动鞋").click()
    p.get_by_role("button", name="改成这些（2 件）").click()
    p.wait_for_function("!location.hash.startsWith('#/outfit')", timeout=20000)   # 保存完会离开挑选页
    shoe = c.item("白色运动鞋")
    assert not shoe.get("worn") and shoe.get("wearsSinceWash") == 0, shoe
    c.go("#/wear")
    expect(p.locator(".garment", has_text="灰色卫衣").first).to_be_visible()


@step("问一问：回答带物品卡片；要修改的先确认再执行")
def _(c):
    p = c.page
    c.go("#/")
    p.locator(".ask-field").click()
    p.get_by_role("button", name="我的充电宝在哪？").click()
    expect(p.locator(".msg.ai").first).to_contain_text("在书桌抽屉")
    expect(p.locator(".msg.ai .row", has_text="手机充电器")).to_be_visible()
    p.locator(".chat-input textarea").fill("牙刷用完了")
    p.locator(".chat-input button").click()
    expect(p.get_by_text("要这样改吗？")).to_be_visible()
    assert c.item("牙刷")["quantity"] == 1, "没确认就改了"
    p.get_by_role("button", name="确认").click()
    expect(p.get_by_text("✓ 已完成 1 项修改")).to_be_visible()
    assert c.item("牙刷")["quantity"] == 0


@step("洗衣篮：今天穿的要洗吗、分批、开洗、收好了；贴身衣物自动收回；床单定期洗")
def _(c):
    p = c.page
    c.go("#/wardrobe")
    p.get_by_text("洗衣篮").click()
    card = p.locator(".card", has_text="要洗吗")
    expect(card).to_contain_text("今天按秋冬的次数默认勾选")                                       # 假天气最高 3°C
    expect(card.locator(".check-row", has_text="灰色卫衣").locator("input")).not_to_be_checked()  # 秋冬上衣 3 次才勾
    expect(card.locator(".check-row", has_text="黑色羽绒服").locator("input")).not_to_be_checked()
    expect(card.locator(".check-row", has_text="运动鞋")).to_have_count(0)                         # 鞋没穿、也不进洗衣篮
    card.locator(".check-row", has_text="灰色卫衣").locator("input").check()
    card.get_by_role("button", name="放进洗衣篮").click()
    expect(p.get_by_text("要洗吗")).to_have_count(0)
    d = c.data()
    assert d["prefs"]["laundryAsked"] and next(i for i in d["items"] if i["id"] == "iw3")["laundry"]["state"] == "dirty"
    sock = next(i for i in d["items"] if i["id"] == "isock")
    assert "laundry" not in sock, f"贴身衣物没自动收回：{sock}"
    expect(p.locator(".group-title", has_text="深色")).to_be_visible()                               # 灰色归深色一批
    expect(p.locator(".section-title", has_text="床上用品该洗了")).to_be_visible()
    c.go(f"#/item/{c.item('黑色羽绒服')['id']}")
    p.get_by_role("button", name="放进洗衣篮").click()
    expect(p.get_by_text("在洗衣篮里")).to_be_visible()
    c.go("#/laundry")
    expect(p.locator(".group-title", has_text="单独洗或送洗")).to_be_visible()                       # 羽绒服单独洗
    p.get_by_role("button", name="开洗勾选的").click()
    expect(p.locator(".section-title", has_text="在洗 / 在晾（2）")).to_be_visible()
    assert not any(i["id"] == "iw3" for i in __import__("json").loads(c.repo.read("inventory.json"))["items"] if not i.get("laundry")), "卫衣应该在洗"
    p.get_by_role("button", name="收好了").click()
    expect(p.locator(".section-title", has_text="在洗 / 在晾")).to_have_count(0)
    w = c.item("灰色卫衣")
    assert "laundry" not in w and w["wearsSinceWash"] == 0 and w["lastWashed"], w
    c.go("#/settings")
    expect(p.locator(".card", has_text="手机提醒")).to_be_visible()


def scan(c, asset):
    c.page.evaluate("code => window.__scan(code)", f"{URL}?a={asset}")


def left_check(c):
    c.page.wait_for_function("!location.hash.includes('/check/')", timeout=20000)


@step("收纳袋：新建、扫码登记里面的东西")
def _(c):
    p = c.page
    c.go("#/boxes")
    p.get_by_role("button", name="新建收纳袋").click()
    sh = p.locator(".sheet")
    sh.get_by_label("名称").fill("洗漱包")
    sh.get_by_label("平时放在哪").select_option(label="储物间 Storage Room")
    sh.get_by_role("button", name="新建").click()
    expect(p.get_by_role("heading", name="洗漱包")).to_be_visible()
    p.get_by_role("link", name="扫码登记袋子里的东西").click()
    expect(p.locator(".check-count")).to_contain_text("0")
    scan(c, c.item("牙刷")["assetId"])
    expect(p.locator(".check-count")).to_contain_text("1")
    p.get_by_role("button", name="保存").click()
    left_check(c)
    bag = next(l for l in c.data()["locations"] if l["name"] == "洗漱包")
    assert bag["box"] == "bag" and bag["parent"] == "Lstore" and bag["assetId"].startswith("010-"), bag
    assert c.item("牙刷")["location"] == bag["id"], c.item("牙刷")
    c.go("#/items?mode=place")
    expect(p.locator(".row.place", has_text="洗漱包")).to_be_visible()   # 收纳袋显示在所在的柜子下面


@step("清单模板：从物品勾选、扫码添加（扫袋子 = 加袋子里的）")
def _(c):
    p = c.page
    c.go("#/lists")
    p.get_by_role("button", name="新建", exact=True).click()
    p.locator(".sheet").get_by_label("名称").fill("回家")
    p.locator(".sheet").get_by_role("button", name="新建").click()
    expect(p.get_by_role("heading", name="回家")).to_be_visible()
    p.get_by_role("button", name="从物品里添加").click()
    p.locator(".sheet input[type=search]").fill("身份证")
    p.locator(".sheet .check-row", has_text="身份证").locator("input").check()
    p.locator(".sheet").get_by_role("button", name="添加").click()
    expect(p.locator(".check-item", has_text="身份证")).to_be_visible()
    p.get_by_role("link", name="扫码添加").click()
    expect(p.locator(".check-count")).to_be_visible()
    scan(c, c.item("手机充电器")["assetId"])
    bag = next(l for l in c.data()["locations"] if l["name"] == "洗漱包")
    scan(c, bag["assetId"])
    expect(p.locator(".check-count")).to_contain_text("3")
    p.get_by_role("button", name="保存模板").click()
    left_check(c)
    expect(p.locator(".check-item")).to_have_count(3)


@step("出行·回家：出发核对（扫码、收纳袋、清单外、不带的）→ 回程核对（落下了、留在家里）→ 找回来了")
def _(c):
    p = c.page
    c.go("#/lists")
    p.locator(".cell", has_text="回家").click()
    p.get_by_role("button", name="用它出行").click()
    expect(p.get_by_role("button", name="回家")).to_have_class("chip on")
    p.get_by_role("button", name="生成清单").click()
    p.get_by_role("button", name="开始出发核对").click()
    expect(p.locator(".check-count")).to_contain_text("/ 3")
    scan(c, c.item("身份证")["assetId"])
    bag = next(l for l in c.data()["locations"] if l["name"] == "洗漱包")
    scan(c, bag["assetId"])                                       # 牙刷在袋子里
    scan(c, c.item("线性代数（第六版）")["assetId"])                 # 清单外 → 问要不要加
    p.locator(".sheet").get_by_role("button", name="加进清单").click()
    expect(p.locator(".check-count")).to_contain_text("3 / 4")
    p.get_by_role("button", name="出发").click()                    # 充电器没确认 → 确认「不带了」
    left_check(c)
    expect(p.get_by_role("link", name="开始回程核对")).to_be_visible()
    d = c.data()
    trip = next(t for t in d["trips"] if t.get("kind") == "回家")
    toothbrush = c.item("牙刷")
    assert sorted(trip["out"]) == sorted([c.item("身份证")["id"], toothbrush["id"], c.item("线性代数（第六版）")["id"]]), trip
    assert toothbrush["homeLocation"] == bag["id"], toothbrush       # 原位置是袋子
    assert c.item("手机充电器")["location"] == "Ldrawer"             # 没带的留在原处
    # 回程
    p.get_by_role("link", name="开始回程核对").click()
    expect(p.locator(".check-count")).to_contain_text("/ 3")
    scan(c, c.item("身份证")["assetId"])
    p.get_by_role("button", name="核对完了").click()
    sh = p.locator(".sheet")
    expect(sh).to_contain_text("还有 2 件没找到")
    sh.locator(".label", has_text="牙刷").get_by_role("button", name="留在家里").click()
    sh.get_by_role("button", name="到了，放回原处").click()
    left_check(c)
    expect(p.get_by_text("出发带了 3 件，回程找到 1 件，落下 1 件")).to_be_visible()
    d = c.data()
    home = next(l for l in d["locations"] if l.get("home"))
    assert c.item("身份证")["location"] == "Ldrawer"
    assert c.item("牙刷")["location"] == home["id"] and "homeLocation" not in c.item("牙刷")
    book = c.item("线性代数（第六版）")
    assert book["location"] == "Ldrawer" and book["leftBehind"]["place"] == "家里", book
    assert not any(l.get("box") == "trip" for l in d["locations"]), "空行李箱没删掉"
    c.go("#/")
    p.locator(".cell", has_text="落在外面了").click()
    p.get_by_role("button", name="找回来了").click()
    expect(p.get_by_text("落在家里了")).to_have_count(0)
    assert "leftBehind" not in c.item("线性代数（第六版）")


@step("搬家箱子：扫码核对拆箱")
def _(c):
    p = c.page
    c.go("#/boxes")
    p.get_by_role("button", name="新建搬家箱子").click()
    p.locator(".sheet").get_by_role("button", name="新建").click()
    expect(p.get_by_role("heading", name="箱子 2")).to_be_visible()
    box = next(l for l in c.data()["locations"] if l.get("box") == "move" and l["name"] != "箱子 1")
    for name in ["手机充电器", "身份证"]:
        c.go(f"#/item/{c.item(name)['id']}/edit")
        p.get_by_label("位置").select_option(label=box["name"])
        p.get_by_role("button", name="保存", exact=True).click()
        c.wait_item(name)
    c.go(f"#/place/{box['id']}")
    p.get_by_role("link", name="核对拆箱").click()
    expect(p.locator(".check-count")).to_be_visible()
    scan(c, c.item("手机充电器")["assetId"])
    expect(p.locator(".check-count")).to_contain_text("1 / 2")
    p.get_by_role("button", name="核对完了").click()
    expect(p.locator(".toast", has_text="还差 1 件：身份证")).to_be_visible()


@step("AI 补全：只填名称，推荐类别和字段")
def _(c):
    p = c.page
    p.evaluate("() => localStorage.setItem('inventory-deepseek', JSON.stringify({ key: 'sk-test' }))")
    c.go("#/new")
    p.reload()
    p.wait_for_function("!localStorage.getItem('inventory-deepseek')")   # 密钥已搬进数据仓库
    # 回归：数据没变时重新打开（用缓存），也要读到数据仓库里的 AI 设置
    p.reload()
    p.wait_for_timeout(1500)
    p.get_by_label("名称").fill("优衣库摇粒绒外套")
    p.get_by_role("button", name="AI 补全").click()
    expect(p.locator(".ai-suggest")).to_contain_text("部位：外套")
    p.locator(".ai-suggest").get_by_role("button", name="采用").click()
    expect(p.get_by_role("button", name="衣服", exact=True)).to_have_class("chip on")
    expect(p.locator(".field-row select").first).to_have_value("外套")
    expect(p.get_by_role("textbox", name="编号")).to_have_value("110-0006")


@step("购物清单：快用完了、手动加、AI 建议、这周不买、勾选、买回来了（补数量、记价格、建档）、剩几件提醒")
def _(c):
    p = c.page
    med = c.item("布洛芬片")
    c.go(f"#/item/{med['id']}")
    p.get_by_role("button", name="快用完了").click()
    expect(p.get_by_text("快用完了（")).to_be_visible()
    c.go("#/shopping")
    expect(p.locator(".shop-row", has_text="布洛芬片")).to_contain_text("快用完了")
    p.get_by_role("button", name="怎么用").click()
    expect(p.locator(".sheet")).to_contain_text("清单里的东西从哪来")
    p.locator(".sheet").get_by_role("button", name="知道了").click()
    ai = p.locator(".shop-ai")
    expect(ai).to_contain_text("护手霜", timeout=20000)
    p.get_by_label("加到购物清单").fill("鸡蛋、5号电池")
    p.get_by_role("button", name="加入", exact=True).first.click()
    expect(p.locator(".shop-row", has_text="5号电池")).to_be_visible()
    ai.get_by_role("button", name="加入").click()
    expect(p.locator(".shop-row", has_text="护手霜")).to_be_visible()
    expect(ai).to_contain_text("这周没有别的要补充")
    p.get_by_role("button", name="5号电池 更多").click()
    p.locator(".sheet").get_by_role("button", name="这周不买").click()
    expect(p.locator(".shop-row", has_text="5号电池")).to_have_count(0)
    for name in ("布洛芬片", "鸡蛋", "护手霜"):
        p.locator(".shop-row", has_text=name).click()
    expect(p.locator(".shop-row.done")).to_have_count(3)
    p.reload()  # 勾选存在手机上，刷新还在
    p.get_by_role("button", name="买回来了（3 样）").click()
    sheet = p.locator(".sheet")
    expect(sheet.get_by_label("布洛芬片 现在有")).to_have_value(str(c.item("布洛芬片")["quantity"] + 1))
    sheet.get_by_label("布洛芬片 现在有").fill("5")
    sheet.get_by_label("布洛芬片 价格").fill("12.5")
    sheet.get_by_label("鸡蛋 价格").fill("8")
    sheet.get_by_label("护手霜 建档").check()
    sheet.get_by_role("button", name="记好了").click()
    # 顺手记账：有价格的两样问要不要记到账本；类别按名字 / 类别猜，可以改；美元账户不出现
    lsheet = p.locator(".sheet", has_text="顺手记一笔账")
    expect(lsheet.get_by_label("布洛芬片 记成")).to_have_value("c-medical")
    expect(lsheet.get_by_label("从哪个账户付").locator("option")).to_have_count(1)
    lsheet.get_by_label("鸡蛋 记成").select_option("c-snack")
    expect(lsheet).to_contain_text("会记 2 笔")
    lsheet.get_by_role("button", name="记到账本").click()
    expect(p.get_by_text("已记到账本")).to_be_visible()
    tx = c.ledger_tx()
    assert sorted((t["category"], t["amount"], t["note"], t["account"]) for t in tx) == [("c-medical", 12.5, "布洛芬片", "a-live"), ("c-snack", 8, "鸡蛋", "a-live")], tx
    expect(p.get_by_text("买回来还没建档")).to_be_visible()
    d = c.data()
    m = next(i for i in d["items"] if i["name"] == "布洛芬片")
    sh = d["shopping"]
    assert m["quantity"] == 5 and "runningLow" not in m, m
    assert [e["name"] for e in sh["extra"]] == ["5号电池"] and "m:" + sh["extra"][0]["id"] in sh["skip"], sh
    assert sorted((x["name"], x.get("price")) for x in sh["history"]) == sorted([("护手霜", None), ("布洛芬片", 12.5), ("鸡蛋", 8)]), sh["history"]
    assert [e["name"] for e in sh["toFile"]] == ["护手霜"], sh
    expect(p.get_by_text("这个月买东西花了 ¥21")).to_be_visible()
    p.get_by_role("link", name="建档").click()
    expect(p.get_by_label("名称")).to_have_value("护手霜")
    p.get_by_label("位置").select_option(label="储物间 Storage Room")
    p.get_by_role("button", name="洗漱护肤", exact=True).click()
    p.get_by_role("button", name="保存", exact=True).click()
    c.wait_item("护手霜")
    assert c.data()["shopping"]["toFile"] == []
    # 能数的消耗品：剩几件提醒
    c.go(f"#/item/{med['id']}/edit")
    p.get_by_label("剩几件提醒").fill("5")
    p.get_by_role("button", name="保存", exact=True).click()
    c.wait_item("布洛芬片")
    assert c.item("布洛芬片")["lowAt"] == 5
    c.go("#/shopping")
    expect(p.locator(".shop-row", has_text="布洛芬片")).to_contain_text("只剩 5")
    c.go("#/stats")
    expect(p.get_by_text("购物花费")).to_be_visible()


@step("顺手记账：新建填了价格问要不要记；小票导入的建档不再问")
def _(c):
    p = c.page
    n = len(c.ledger_tx())
    c.go("#/new")
    p.get_by_label("名称").fill("充电头")
    p.get_by_label("位置").select_option(label="　　书桌抽屉 Desk Drawer")
    p.get_by_role("button", name="电子产品", exact=True).click()
    p.get_by_text("品牌、购买与保修").click()
    p.get_by_label("价格（元）").fill("39")
    p.get_by_role("button", name="保存", exact=True).click()
    c.wait_item("充电头")
    lsheet = p.locator(".sheet", has_text="顺手记一笔账")
    expect(lsheet.get_by_label("充电头 记成")).to_have_value("c-gadget")
    lsheet.get_by_role("button", name="记到账本").click()
    expect(p.get_by_text("已记到账本")).to_be_visible()
    tx = c.ledger_tx()
    assert len(tx) == n + 1 and tx[-1]["note"] == "充电头" and tx[-1]["amount"] == 39 and tx[-1]["date"] == date.today().isoformat(), tx[-1]
    # 账本的小票导入放进「还没建档」的（已经记过账）：建档时带上价格、数量，不再问记账
    d = c.data()
    d.setdefault("shopping", {}).setdefault("toFile", []).append({"id": "mrcpt", "name": "收纳盒", "price": 25, "qty": 2, "date": date.today().isoformat(), "from": "receipt", "paid": True})
    c.repo.external_write("inventory.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/")
    p.reload()
    expect(p.locator(".cell", has_text="1 样买回来还没建档")).to_be_visible()
    c.go("#/shopping")
    expect(p.get_by_text("（小票导入）")).to_be_visible()
    p.get_by_role("link", name="建档").click()
    expect(p.get_by_label("名称")).to_have_value("收纳盒")
    expect(p.get_by_label("价格（元）")).to_have_value("25")
    expect(p.get_by_label("数量")).to_have_value("2")
    p.get_by_label("位置").select_option(label="储物间 Storage Room")
    p.get_by_role("button", name="日用杂物", exact=True).click()
    p.get_by_role("button", name="保存", exact=True).click()
    c.wait_item("收纳盒")
    p.wait_for_timeout(500)
    expect(p.locator(".sheet", has_text="顺手记一笔账")).to_have_count(0)
    assert len(c.ledger_tx()) == n + 1 and c.data()["shopping"]["toFile"] == []


@step("用完了、重新打印直接做，可以撤销；没网也能改，有网自动上传并和另一台设备的修改合并")
def _(c):
    p = c.page
    tb = c.item("布洛芬片")
    c.go(f"#/item/{tb['id']}")
    p.get_by_role("button", name="用完了", exact=True).click()
    toast = p.locator(".toast.undo")
    expect(toast).to_contain_text("用完了，已放进购物清单")
    assert c.item("布洛芬片")["quantity"] == 0
    toast.get_by_role("button", name="撤销").click()
    expect(p.get_by_text("已撤销")).to_be_visible()
    t = c.item("布洛芬片")
    assert t["quantity"] == tb["quantity"] and t.get("notes", "") == tb.get("notes", ""), t
    # 没网：先存手机，页面马上变，顶上提示；刷新还在；有网后自动上传，别的设备的修改不丢
    p.route(f"{API}/**", lambda r: r.abort())
    p.get_by_role("button", name="快用完了").click()
    expect(p.get_by_text("快用完了（")).to_be_visible()
    expect(p.locator(".busy")).to_have_count(0)
    expect(p.locator(".sync-pill")).to_contain_text("没网，1 项存在手机上")
    p.reload()
    expect(p.get_by_text("快用完了（")).to_be_visible()
    assert "runningLow" not in json.loads(c.repo.read("inventory.json"))["items"][[i["id"] for i in json.loads(c.repo.read("inventory.json"))["items"]].index(tb["id"])]
    d = json.loads(c.repo.read("inventory.json"))
    next(i for i in d["items"] if i["name"] == "身份证")["notes"] = "另一台设备写的"
    c.repo.external_write("inventory.json", json.dumps(d, ensure_ascii=False).encode())
    p.unroute(f"{API}/**")
    p.evaluate("window.dispatchEvent(new Event('online'))")
    d = c.data()
    assert next(i for i in d["items"] if i["id"] == tb["id"]).get("runningLow"), "离线的修改没传上去"
    assert next(i for i in d["items"] if i["name"] == "身份证")["notes"] == "另一台设备写的"
    expect(p.locator(".sync-pill")).to_be_hidden()


@step("换季整理：列出该收起来的，按勾选移动")
def _(c):
    p = c.page
    c.go("#/season")
    expect(p.locator(".section-title", has_text="收起来")).to_be_visible()
    before = {i["id"]: i["location"] for i in c.data()["items"]}
    p.get_by_role("button", name="按勾选整理").click()
    expect(p.get_by_text("当季衣柜里正好都是该穿的")).to_be_visible()
    after = {i["id"]: i["location"] for i in c.data()["items"]}
    assert after != before, "没有移动任何衣服"


@step("设置：常住城市存进数据")
def _(c):
    p = c.page
    c.go("#/settings")
    card = p.locator(".card", has_text="常住城市")
    card.locator("input").fill("北京")
    card.get_by_role("button", name="保存").click()
    # 上一步的「已保存」提示可能还没消失，所以直接等数据写进去
    for _ in range(50):
        if c.data()["prefs"].get("homeCity") == "北京":
            break
        p.wait_for_timeout(200)
    assert c.data()["prefs"]["homeCity"] == "北京"


@step("导出全部物品 Excel")
def _(c):
    import openpyxl
    p = c.page
    c.go("#/me")
    with p.expect_download() as dl:
        p.get_by_role("button", name="导出全部物品（Excel）").click()
    path = ART / "all.xlsx"
    dl.value.save_as(path)
    rows = list(openpyxl.load_workbook(path).active.iter_rows(values_only=True))
    assert rows[0][:3] == ("编号", "名称", "类别") and len(rows) == len(c.data()["items"]) + 1


@step("批量录入脚本：自动编号接着网页的往后排")
def _(c):
    manifest = ART / "manifest.json"
    manifest.write_text(json.dumps({"location": "书桌抽屉", "items": [
        {"name": "数据线", "tags": ["电子产品"], "quantity": 2},
        {"name": "条纹衬衫", "tags": ["衣服"], "location": "当季衣柜", "fields": {"季节": "春秋"}},
        {"name": "薯片", "tags": ["零食食品"], "location": "零食食品柜"},
    ]}, ensure_ascii=False))
    env = {**os.environ, "GITHUB_API_URL": API, "GITHUB_TOKEN": "test-token", "NO_PROXY": "127.0.0.1,localhost", "no_proxy": "127.0.0.1,localhost"}
    out = subprocess.run([sys.executable, str(ROOT / "tools/import_items.py"), "--repo", REPO, str(manifest)],
                         capture_output=True, text=True, env=env)
    assert out.returncode == 0, out.stdout + out.stderr
    d = c.data()
    got = {i["name"]: (i["assetId"], i["label"], i["consumable"]) for i in d["items"] if i["name"] in ("数据线", "条纹衬衫", "薯片")}
    assert got == {"数据线": ("100-0006", "pending", False), "条纹衬衫": ("110-0006", "none", False),
                   "薯片": ("280-0001", "pending", True)}, got


# ---------- 假的天气和 DeepSeek ----------

def fake_externals(page):
    page.route("https://geocoding-api.open-meteo.com/**", lambda r: r.fulfill(json={
        "results": [{"name": "上海", "admin1": "上海", "latitude": 31.2, "longitude": 121.5}]}))
    page.route("https://api.open-meteo.com/**", lambda r: r.fulfill(json={
        "hourly": {"temperature_2m": [8] * 7 + [9, 10, 11, 13, 14, 16, 17, 16, 15, 14, 13, 12, 11, 10, 10, 9],
                   "precipitation_probability": [10] * 24, "wind_speed_10m": [8] * 24},
        "daily": {"time": [D(1), D(2), D(3), D(4)], "temperature_2m_min": [-5, 6, 7, 8], "temperature_2m_max": [3, 15, 16, 14],
                  "precipitation_probability_max": [10, 60, 20, 10], "weather_code": [71, 61, 3, 3]}}))
    trip = {"items": [{"id": "iw1", "qty": 1, "reason": "早晚冷"}, {"id": "iw3", "qty": 1, "reason": "白天穿"},
                      {"id": "iid", "qty": 1, "reason": "必带"}, {"id": "不存在的东西", "qty": 1, "reason": "AI 编的"}],
            "outfits": [{"day": D(2)[5:], "items": ["iw3", "iw1"], "note": "灰配黑"}],
            "missing": [{"name": "转换插头", "reason": "酒店插座"}], "tips": ["带伞"]}
    outfit = {"options": [{"title": "灰卫衣配黑羽绒", "items": ["iw3", "iw1", "ish", "不存在"], "why": "早上冷", "tips": ["中午可以脱外套"]},
                          {"title": "暖和的一套", "items": ["iw1", "ish"], "why": "更暖", "tips": []}]}
    autofill = {"category": "衣服", "fields": {"部位": "外套", "季节": "冬", "厚薄": "厚", "颜色": "灰"}, "consumable": False, "note": ""}

    def chat(user):
        if "充电宝" in user:
            return {"answer": "在书桌抽屉里。", "items": ["ich"], "actions": []}
        if "牙刷" in user:
            return {"answer": "要把牙刷标记为用完吗？", "items": ["itb"], "actions": [{"type": "use_up", "id": "itb"}]}
        return {"answer": "不知道", "items": [], "actions": []}

    def deepseek(route):
        if route.request.url.endswith("/models"):
            return route.fulfill(json={"data": [{"id": "deepseek-flash"}, {"id": "deepseek-v4-pro"}]})
        body = json.loads(route.request.post_data)
        system, user = body["messages"][0]["content"], body["messages"][-1]["content"]
        if "收拾行李" in system:
            ans = trip
        elif "搭配日常穿着" in system:
            ans = outfit
        elif "去超市" in system:
            ans = {"suggestions": [{"name": "护手霜", "reason": "降温了，手容易干"}]}
        elif "给宿舍里的物品建档" in system:
            ans = autofill
        else:
            ans = chat(user)
        route.fulfill(json={"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(ans, ensure_ascii=False)}}]})
    page.route("https://api.deepseek.com/**", deepseek)


@step("找东西（Siri）：「充电器在哪」去掉在哪找名字，按字也能找；Siri 说明页")
def _(c):
    p = c.page
    c.go("#/find?text=" + quote("手机充电器在哪"))
    hit = p.locator(".find-hit", has_text="手机充电器")
    expect(hit).to_contain_text("书桌抽屉")
    c.go("#/find?text=" + quote("充电线放哪了"))  # 名字里没有「充电线」，按字找到「手机充电器」
    expect(p.locator(".find-hit", has_text="手机充电器")).to_be_visible()
    c.go("#/find?text=" + quote("游泳镜在哪里"))
    expect(p.get_by_text("档案里没找到")).to_be_visible()
    c.go("#/siri")
    expect(p.locator("code", has_text="#/find?text=")).to_be_visible()


@step("退货期：新买的电子产品自动填退货截止，快到了首页提醒，「没问题，不退了」")
def _(c):
    p = c.page
    c.go("#/new")
    p.get_by_label("名称").fill("蓝牙耳机")
    p.get_by_label("位置").select_option(label="　　书桌抽屉 Desk Drawer")
    p.get_by_role("button", name="电子产品", exact=True).click()
    p.get_by_text("品牌、购买与保修").click()
    p.get_by_label("购买日期", exact=True).fill(D(-5))
    p.get_by_role("button", name="保存", exact=True).click()
    c.wait_item("蓝牙耳机")
    if p.locator(".sheet", has_text="顺手记一笔账").count():
        p.locator(".sheet").get_by_role("button", name="不用了").click()
    assert c.item("蓝牙耳机")["returnBy"] == D(2)
    expect(p.get_by_text(f"{D(2)} 前还能退")).to_be_visible()
    c.go("#/")
    expect(p.locator(".cell", has_text="蓝牙耳机 2 天后过退货期")).to_be_visible()
    p.locator(".cell", has_text="蓝牙耳机 2 天后过退货期").click()
    p.get_by_role("button", name="没问题，不退了").click()
    expect(p.get_by_text("前还能退")).to_have_count(0)
    assert "returnBy" not in c.item("蓝牙耳机")


@step("消耗品多久买一次、衣服穿一次多少钱、购物清单估价和预算、账本里还剩多少")
def _(c):
    p = c.page
    d = c.data()
    med = next(i for i in d["items"] if i["name"] == "布洛芬片")
    d.setdefault("shopping", {}).setdefault("history", [])
    d["shopping"]["history"] = [x for x in d["shopping"]["history"] if x.get("itemId") != med["id"]] + [{"date": D(-30), "name": "布洛芬片", "itemId": med["id"], "price": 12},
                                 {"date": D(-15), "name": "布洛芬片", "itemId": med["id"], "price": 13}]
    med.pop("runningLow", None); med.pop("lowAt", None)
    d["shopping"]["extra"], d["shopping"]["skip"] = [], {}
    coat = next(i for i in d["items"] if i["name"] == "黑色羽绒服")
    coat.update(purchasePrice=600, worn=[D(-3), D(-2), D(-1)])
    tee = next(i for i in d["items"] if i["name"] == "灰色卫衣")
    tee.update(purchasePrice=200, worn=[])
    c.repo.external_write("inventory.json", json.dumps(d, ensure_ascii=False).encode())
    c.go(f"#/item/{med['id']}")
    p.reload()
    expect(p.get_by_text("平均 15 天")).to_be_visible()
    c.go(f"#/item/{coat['id']}")
    expect(p.get_by_text("¥200（穿了 3 次）")).to_be_visible()
    c.go("#/shopping")
    expect(p.locator(".shop-row", has_text="布洛芬片")).to_contain_text("平时 15 天买一次")
    expect(p.locator(".shop-budget")).to_contain_text("大概 ¥13")
    expect(p.locator(".shop-budget")).to_contain_text("账本：这个月日常还剩")
    c.prompt = "10"
    p.get_by_role("button", name="设预算").click()
    expect(p.locator(".shop-budget")).to_contain_text("比预算多 ¥3")
    assert c.data()["prefs"]["shopBudget"] == 10
    c.go("#/stats")
    expect(p.get_by_text("消耗品每月大约")).to_be_visible()
    cpw = p.locator(".card", has_text="衣服穿一次多少钱")
    expect(cpw.locator(".cpw-row", has_text="黑色羽绒服")).to_contain_text("¥200/次")
    expect(cpw.locator(".cpw-row", has_text="灰色卫衣")).to_contain_text("还没穿")


@step("放假离校清单：会过期的、贵重东西、每次要做的；首页提醒；打勾；建成回家出行")
def _(c):
    p = c.page
    d = c.data()
    d["items"].append({**c.item("布洛芬片"), "id": "isnack", "name": "牛肉干", "assetId": None, "tags": ["零食食品"], "quantity": 2,
                       "fields": {"保质期": D(20)}, "notes": ""})
    c.repo.external_write("inventory.json", json.dumps(d, ensure_ascii=False).encode())
    c.go("#/term")
    p.reload()
    p.get_by_label("返校日期").fill(D(40))
    p.get_by_label("返校日期").dispatch_event("change")
    expect(p.get_by_label("返校日期")).to_have_value(D(40))
    p.get_by_label("离校日期").fill(D(3))
    p.get_by_label("离校日期").dispatch_event("change")
    expect(p.locator(".shop-row", has_text="牛肉干")).to_contain_text(f"保质期 {D(20)}")
    expect(p.locator(".shop-row", has_text="身份证")).to_be_visible()
    p.locator(".shop-row", has_text="拔掉插头").click()
    expect(p.locator(".shop-row.done", has_text="拔掉插头")).to_be_visible()
    t = c.data()["term"]
    assert t["leave"] == D(3) and t["back"] == D(40) and any(k.startswith("task:leave") for k in t["done"]), t
    c.go("#/")
    expect(p.locator(".cell", has_text="3 天后离校：清单还有")).to_be_visible()
    c.go("#/term")
    p.get_by_role("button", name="建成回家出行，扫码装包").click()
    p.wait_for_function("location.hash === '#/trip/new'")
    assert any(l["name"] == "放假带回家" and c.item("身份证")["id"] in l["items"] for l in c.data()["lists"])
    c.go("#/term?kind=back")
    expect(p.locator(".shop-row", has_text="开窗通风")).to_be_visible()


@step("手机丢了怎么办：怎么删令牌，最近的修改记录带设备名")
def _(c):
    p = c.page
    c.go("#/settings")
    p.get_by_role("link", name="手机丢了怎么办").click()
    expect(p.get_by_text("最下面 Delete")).to_be_visible()
    expect(p.get_by_text("这台是")).to_be_visible()
    expect(p.locator(".commit-row").first).to_be_visible()
    assert " · " in c.repo.commits[c.repo.head]["message"]  # 网页提交的说明带上了设备名


@step("今天穿什么分五类：运动只在跑步、居家只在宅宿舍、休闲爬山、正式上班、隆重只在隆重场合；13 度以下提醒秋衣；腰带累了提醒换")
def _(c):
    p = c.page
    c.go("#/")
    out = p.evaluate("""async () => {
      const { ruleOutfit, styleOf, dayStyles } = await import('./js/outfit.js');
      const I = (id, name, part, style, tag = '衣服') => ({ id, name, tags: [tag], fields: { 部位: part, 季节: '春秋', ...(style ? { 风格: style } : {}) } });
      const data = { locations: [], items: [
        I('f1', '黑色衬衫', '上衣', '正式'), I('f2', '黑色西裤', '下装', ''), I('f3', '皮鞋', '鞋', '正式', '鞋'),
        I('c1', '格子衬衫', '上衣', '休闲'), I('c2', '工装裤', '下装', '休闲'), I('c3', '登山鞋', '鞋', '休闲', '鞋'),
        I('r1', '蓝色长袖运动上衣', '上衣', '', '运动服'), I('r2', '黑色运动裤子', '下装', ''), I('r3', '白色运动鞋', '鞋', '', '鞋'),
        I('h1', '睡衣', '上衣', '休闲'), I('h2', '珊瑚绒裤', '下装', '居家'), I('l1', '秋衣', '上衣', ''),
        I('g1', '黑色西装上衣', '上衣', ''), I('g2', '西装裤', '下装', '隆重'),
      ] };
      const w = { min: 12, max: 20 };
      const plan = (s) => ruleOutfit(data, w, s).options.map((o) => ({ title: o.title, items: o.items.sort(), tips: o.tips }));
      return { styles: data.items.map((i) => styleOf(i)), work: plan(['上班']), hike: plan(['爬山 / 出去玩']), run: plan(['上班', '跑步']),
               home: plan(['宅宿舍']), grand: plan(['隆重场合']), def: dayStyles([]) };
    }""")
    assert out["styles"] == ["正式", "正式", "正式", "休闲", "休闲", "休闲", "运动", "运动", "运动", "居家", "居家", "正式", "隆重", "隆重"], out["styles"]
    assert out["grand"][0]["items"] == ["f3", "g1", "g2"], out["grand"]  # 西装 + 正式的皮鞋
    assert out["work"][0]["items"] == ["f1", "f2", "f3"] and len(out["work"]) == 1, out["work"]
    assert any("秋衣" in t for t in out["work"][0]["tips"]), "最低 12 度（低于 13）提醒加秋衣"
    assert out["hike"][0]["items"] == ["c1", "c2", "c3"], out["hike"]
    assert out["run"][0]["items"] == ["f1", "f2", "f3"] and out["run"][1]["items"] == ["r1", "r2", "r3"], out["run"]
    assert [o["items"] for o in out["home"]] == [["h1", "h2"]], out["home"]
    assert out["def"]["main"] == ["正式"]
    # 腰带：系一天 +1、歇一天 −0.5；连着系满 7 天提醒换成歇着的那条
    days = lambda a, b: [D(-n) for n in range(a, b - 1, -1)]  # noqa: E731
    f = p.evaluate("""async (args) => {
      const { beltFatigue, beltAdvice } = await import('./js/outfit.js');
      const today = args.today;
      const B = (id, name, worn) => ({ id, name, tags: ['衣服'], fields: { 部位: '配饰' }, worn });
      const tired = B('b1', '黑色腰带', args.week), rested = B('b2', '棕色腰带', args.old);
      const alt = B('b3', '轮着系的腰带', args.alt);
      return { tired: beltFatigue(tired, today), alt: beltFatigue(alt, today), rested: beltFatigue(rested, today),
               adv: beltAdvice({ items: [tired, rested] }, today) };
    }""", {"today": D(0), "week": days(6, 0), "old": days(20, 14), "alt": [D(-n) for n in range(0, 20, 2)]})
    assert f["tired"]["score"] == 7 and f["tired"]["streak"] == 7, f["tired"]
    assert f["rested"]["score"] == 0.5 and f["rested"]["rested"] == 14, f["rested"]  # 系 7 天累到 7，歇 13 天剩 0.5，今天过完就回到 0
    assert f["alt"]["score"] < 7, f["alt"]
    assert f["adv"]["swap"]["b"]["id"] == "b2", f["adv"]
    d = c.data()
    belt = next(i for i in d["items"] if i["name"] == "身份证")  # 借一个物品改成两条腰带
    d["items"] += [{**belt, "id": "ibelt1", "name": "黑色腰带", "tags": ["衣服"], "fields": {"部位": "配饰"}, "worn": days(6, 1), "label": "none"},
                   {**belt, "id": "ibelt2", "name": "棕色腰带", "tags": ["衣服"], "fields": {"部位": "配饰"}, "worn": [D(-30)], "label": "none"}]
    c.repo.external_write("inventory.json", json.dumps(d, ensure_ascii=False).encode())
    # 厚薄跟着季节：夏季只有薄 / 厚；气温决定最合适的「季节·厚薄」
    fit = p.evaluate("""async () => {
      const { thickOptions, fitScore } = await import('./js/outfit.js');
      const I = (season, thick) => ({ fields: { 季节: season, 厚薄: thick } });
      return { summer: thickOptions('夏'), winter: thickOptions('冬'),
               hot: [fitScore(I('夏', '薄'), 30), fitScore(I('夏', '厚'), 30)],
               cold: [fitScore(I('冬', '厚'), 0), fitScore(I('冬', '适中'), 0), fitScore(I('冬', '薄'), 0), fitScore(I('夏', '薄'), 0)],
               mild: [fitScore(I('冬', '适中'), 6), fitScore(I('冬', '厚'), 6)] };
    }""")
    assert fit["summer"] == ["薄", "厚"] and fit["winter"] == ["薄", "适中", "厚"], fit
    assert fit["hot"][0] > fit["hot"][1] and fit["cold"][0] > fit["cold"][1] > fit["cold"][2] > fit["cold"][3] and fit["mild"][0] > fit["mild"][1], fit
    # 新建衣服：没有「＋尺码」；季节选夏，厚薄只能选薄 / 厚
    c.go("#/new")
    p.get_by_role("button", name="衣服", exact=True).click()
    expect(p.get_by_role("button", name="＋颜色")).to_be_visible()
    expect(p.get_by_role("button", name="＋尺码")).to_have_count(0)
    p.get_by_role("button", name="＋季节").click()
    p.get_by_label("季节", exact=True).select_option("夏")
    p.get_by_role("button", name="＋厚薄").click()
    expect(p.get_by_label("厚薄", exact=True).locator("option")).to_have_text(["选择…", "薄", "厚"])
    c.go("#/outfit?pick=1")
    p.reload()
    expect(p.locator(".belt-banner")).to_contain_text("「黑色腰带」已经连着系了 6 天")
    expect(p.locator(".belt-line")).to_contain_text("黑色腰带 累 6/7")
    # 同一种腰带两条（数量 ×2）：按「正在系的 / 歇着的」算，换好了点一下
    d = c.data()
    d["items"] = [i for i in d["items"] if i["id"] not in ("ibelt1", "ibelt2")]
    d["items"].append({**belt, "id": "ibelt", "name": "腰带", "tags": ["衣服"], "quantity": 2, "fields": {"部位": "配饰"},
                       "worn": [D(-20)] + days(6, 1), "beltSwaps": [{"date": D(-21), "to": 1}, {"date": D(-19), "to": 0}], "label": "none"})
    c.repo.external_write("inventory.json", json.dumps(d, ensure_ascii=False).encode())
    p.reload()
    expect(p.locator(".belt-banner")).to_contain_text("「腰带」正在系的这条已经连着系了 6 天")
    expect(p.locator(".belt-line")).to_contain_text("腰带（正在系的） 累 6/7")
    p.locator(".belt-banner").get_by_role("button", name="换好了").click()
    expect(p.get_by_text("今天起系的是另一条")).to_be_visible()
    expect(p.locator(".belt-banner")).to_have_count(0)
    expect(p.locator(".belt-line")).to_contain_text("腰带（歇着的） 累 6/7")
    assert c.item("腰带")["beltSwaps"][-1] == {"date": D(0), "to": 1}


@step("序列号：扫条形码填进去（去掉 S/N 前缀），取消不改")
def _(c):
    p = c.page
    c.go("#/new")
    p.get_by_label("名称").fill("移动硬盘")
    p.get_by_label("位置").select_option(label="　　书桌抽屉 Desk Drawer")
    p.get_by_role("button", name="电子产品", exact=True).click()
    p.get_by_text("品牌、购买与保修").click()
    p.get_by_role("button", name="扫条形码").click()
    sheet = p.locator(".sheet", has_text="扫序列号")
    expect(sheet.locator(".scan-frame.bar")).to_be_visible()
    p.wait_for_function("() => typeof window.__scan === 'function'")
    p.evaluate("() => window.__scan('S/N: WX12A3456789')")
    expect(sheet).to_have_count(0)
    expect(p.get_by_label("序列号")).to_have_value("WX12A3456789")
    # 再扫一次但取消：原来的不变
    p.get_by_role("button", name="扫条形码").click()
    sheet.get_by_role("button", name="取消").click()
    expect(sheet).to_have_count(0)
    assert p.evaluate("() => window.__scan === undefined")
    expect(p.get_by_label("序列号")).to_have_value("WX12A3456789")
    p.get_by_role("button", name="保存", exact=True).click()
    c.wait_item("移动硬盘")
    assert c.item("移动硬盘")["serialNumber"] == "WX12A3456789"
    # 顺手记账的弹窗（没填价格不会问）
    expect(p.locator(".sheet", has_text="顺手记一笔账")).to_have_count(0)


def main():
    only = sys.argv[1:]
    ART.mkdir(exist_ok=True)
    repo = FakeRepo(seed())
    ledger = FakeRepo(ledger_seed())
    serve({REPO: repo, "test/finance-data": ledger}, API_PORT)
    handler = partial(SimpleHTTPRequestHandler, directory=str(ROOT))
    handler.log_message = lambda *a: None
    app = ThreadingHTTPServer(("127.0.0.1", APP_PORT), handler)
    threading.Thread(target=app.serve_forever, daemon=True).start()

    failed, errors, ran = [], [], 0
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        ctx = browser.new_context(viewport={"width": 390, "height": 844}, accept_downloads=True)
        ctx.set_default_timeout(10000)
        page = ctx.new_page()
        page.on("pageerror", lambda e: errors.append(str(e)))
        c = Ctx(page, repo, ledger)
        page.on("dialog", lambda d: d.accept(c.prompt) if d.type == "prompt" else d.accept())
        fake_externals(page)
        for i, (name, fn) in enumerate(STEPS):
            if i and only and not any(k in name for k in only):
                continue
            ran += 1
            try:
                fn(c)
                print(f"  ✓ {name}")
            except Exception:
                failed.append(name)
                page.screenshot(path=ART / f"fail-{i:02d}.png", full_page=True)
                print(f"  ✗ {name}\n{traceback.format_exc()}")
                if i == 0:
                    break
        browser.close()
    if errors:
        print("页面报错：", *errors, sep="\n  ")
    print(f"\n{ran - len(failed)}/{ran} 通过" + ("" if ran == len(STEPS) else f"（共 {len(STEPS)} 项，没跑完）"))
    sys.exit(1 if failed or errors else 0)


if __name__ == "__main__":
    main()
